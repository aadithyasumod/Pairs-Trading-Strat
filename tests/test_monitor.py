import json
import tempfile
import unittest
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from pairs_trading.data import generate_synthetic_universe
from pairs_trading.simple_monitor import _load_config, _run_once, _save_run_excel, _write_minimal_xlsx


class MonitorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        prices = generate_synthetic_universe(n_days=600, seed=7, end="2026-09-18").prices
        prices.reset_index().to_csv(self.dir / "prices.csv", index=False)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def write_config(self, **overrides) -> Path:
        cfg = {
            "csv": "prices.csv",
            "min_corr": 0.5,
            "max_data_age_days": None,
            "bell_alert": False,
            "out_csv": "out/latest.csv",
            "out_json": "out/summary.json",
            "results_dir": "out/results",
        }
        cfg.update(overrides)
        path = self.dir / "monitor.json"
        path.write_text(json.dumps(cfg))
        return path

    def test_relative_paths_resolve_against_the_config_file(self) -> None:
        cfg = _load_config(str(self.write_config()))
        self.assertEqual(Path(cfg.csv), (self.dir / "prices.csv").resolve())
        self.assertEqual(Path(cfg.results_dir), (self.dir / "out/results").resolve())

    def test_unknown_config_keys_are_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "Unknown keys"):
            _load_config(str(self.write_config(z_entery=2.0)))

    def test_run_once_ranks_the_cointegrated_pairs(self) -> None:
        cfg = _load_config(str(self.write_config()))
        df, summary = _run_once(cfg)
        found = set(zip(df["asset_a"], df["asset_b"]))
        # Only planted pairs may appear; with a 252-bar screen not every one reaches p <= 0.05.
        self.assertTrue(found <= {("P1_A", "P1_B"), ("P2_A", "P2_B"), ("P3_A", "P3_B")}, found)
        self.assertGreaterEqual(len(found), 1)
        self.assertEqual(summary["price_data_as_of"], "2026-09-18T00:00:00")
        self.assertIn(summary["strategy_traffic_light"], {"GREEN", "YELLOW", "RED"})
        for col in ("deployable_rank", "is_profitable_candidate", "candidate_reasons", "formation_half_life"):
            self.assertIn(col, df.columns)

    def test_stale_data_is_rejected(self) -> None:
        old = generate_synthetic_universe(n_days=600, seed=7, end="2025-01-03").prices
        old.reset_index().to_csv(self.dir / "old_prices.csv", index=False)
        cfg = _load_config(str(self.write_config(csv="old_prices.csv", max_data_age_days=7)))
        with self.assertRaisesRegex(ValueError, "stale"):
            _run_once(cfg)

    def test_run_workbook_is_numbered_and_readable(self) -> None:
        cfg = _load_config(str(self.write_config()))
        df, summary = _run_once(cfg)
        first = _save_run_excel(cfg, df, summary)
        second = _save_run_excel(cfg, df, summary)
        self.assertEqual((first.name, second.name), ("run_0001.xlsx", "run_0002.xlsx"))
        wb = openpyxl.load_workbook(second)
        self.assertEqual(wb.sheetnames, ["opportunities", "summary"])

    def test_fallback_writer_produces_a_valid_workbook(self) -> None:
        df = pd.DataFrame(
            {
                "name": ["A", "B&C"],
                "flag": [True, False],
                "value": [1.5, np.nan],
                "when": pd.to_datetime(["2026-09-18", "2026-09-19"]),
            }
        )
        path = self.dir / "fallback.xlsx"
        _write_minimal_xlsx(path, {"sheet": df})
        ws = openpyxl.load_workbook(path)["sheet"]
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("name", "flag", "value", "when"))
        self.assertEqual(rows[1][:3], ("A", True, 1.5))
        self.assertEqual(rows[2][:3], ("B&C", False, None))


if __name__ == "__main__":
    unittest.main()
